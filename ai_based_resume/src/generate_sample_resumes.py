"""
Sample Resume Generator
Generates realistic multi-domain candidate resumes in TXT, DOCX, and PDF formats.
Domains:
- NLP & Generative AI Engineer
- Java Spring Boot Backend Engineer
- Big Data & Data Engineer
- Cyber Security Analyst
- Frontend React / Next.js Engineer
- Cloud Solutions Architect
- Mobile Developer (iOS / Android / Flutter)
- BI & Business Data Analyst
"""

import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

OUTPUT_DIR = "data/sample_resumes"
os.makedirs(OUTPUT_DIR, exist_ok=True)

RESUMES_DATA = [
    {
        "filename_base": "emily_chen_nlp_genai",
        "name": "Emily Chen",
        "title": "Lead NLP & Generative AI Engineer",
        "contact": "emily.chen@ai-innovate.com | (555) 312-8901 | Seattle, WA",
        "links": "LinkedIn: linkedin.com/in/emilychen-nlp | GitHub: github.com/emilychen-ai",
        "summary": "Innovative NLP and Generative AI Engineer with 5+ years of experience building production Transformer models, Large Language Model (LLM) fine-tuning pipelines, and RAG architectures. Deep expertise with PyTorch, HuggingFace, LangChain, spaCy, and vector databases.",
        "experience": [
            "Lead AI Engineer | GenAI Labs (2022 - Present)\n- Spearheaded enterprise Retrieval-Augmented Generation (RAG) platform using LangChain, HuggingFace Transformers, and Pinecone, handling 4M queries/day.\n- Fine-tuned open-source LLMs (Llama 3, Mistral) with PEFT/LoRA using PyTorch on AWS EC2 GPU clusters.\n- Implemented real-time semantic entity extraction and sentiment analysis pipelines with spaCy and FastAPI.",
            "Senior NLP Scientist | Cognition Works (2019 - 2022)\n- Designed multilingual text classification and summarization algorithms using BERT and Scikit-Learn.\n- Cleaned and prepared 100M+ unstructured document corpora using Python, Pandas, and SQL.\n- Containerized machine learning services with Docker and deployed onto AWS EKS via GitHub Actions."
        ],
        "education": "Master of Science in Artificial Intelligence, Stanford University (2017 - 2019)\nBachelor of Science in Computer Science, UC Berkeley (2013 - 2017)",
        "skills": "Python, PyTorch, HuggingFace, Transformers, LLMs, LangChain, RAG, NLP, spaCy, NLTK, Scikit-Learn, Docker, AWS, Git, SQL, FastAPI, System Design, Problem Solving",
        "formats": ["txt", "pdf"]
    },
    {
        "filename_base": "david_kim_java_backend",
        "name": "David Kim",
        "title": "Senior Java & Microservices Backend Engineer",
        "contact": "david.kim@backend-systems.net | (555) 421-9876 | New York, NY",
        "links": "LinkedIn: linkedin.com/in/davidkim-java | GitHub: github.com/davidkim-dev",
        "summary": "Backend Software Architect with 6+ years of experience engineering high-throughput distributed systems, microservices, and asynchronous event streams. Expert in Java, Spring Boot, Apache Kafka, PostgreSQL, Docker, and AWS.",
        "experience": [
            "Senior Backend Engineer | FinStream Global (2021 - Present)\n- Engineered resilient fintech microservices using Java 21, Spring Boot, and Spring Cloud, processing 15,000 transactions/sec.\n- Designed event-driven messaging pipelines utilizing Apache Kafka and RabbitMQ with zero message loss.\n- Optimized relational database performance in PostgreSQL with Redis caching, reducing p99 latency by 55%.\n- Integrated continuous integration and deployment with Jenkins, Docker, and Kubernetes on AWS.",
            "Java Software Engineer | OmniCommerce (2018 - 2021)\n- Developed core RESTful APIs and payment processing gateways using Java, Spring MVC, and Hibernate.\n- Automated unit and integration testing suites with JUnit, Mockito, and Postman, achieving 90%+ test coverage."
        ],
        "education": "Bachelor of Science in Software Engineering, Cornell University (2014 - 2018)",
        "skills": "Java, Spring Boot, Microservices, REST API, Kafka, PostgreSQL, MySQL, Redis, Docker, Kubernetes, AWS, Git, CI/CD, Unit Testing, Agile, OOP",
        "formats": ["txt", "docx"]
    },
    {
        "filename_base": "priya_sharma_data_engineer",
        "name": "Priya Sharma",
        "title": "Senior Big Data & ETL Engineer",
        "contact": "priya.sharma@datatech.org | (555) 765-4321 | San Jose, CA",
        "links": "LinkedIn: linkedin.com/in/priyasharma-data | GitHub: github.com/priyasharma",
        "summary": "Data Engineer with 4+ years of expertise designing modern data lakehouses, robust ETL/ELT pipelines, and scalable distributed data architectures. Proficient in Apache Spark, PySpark, Airflow, Snowflake, SQL, and GCP BigQuery.",
        "experience": [
            "Senior Data Engineer | DataSphere Corp (2022 - Present)\n- Architected petabyte-scale data pipelines using PySpark and Apache Spark on Databricks clusters.\n- Orchestrated over 200 automated workflows and data quality checks using Apache Airflow.\n- Migrated enterprise data warehouse from legacy SQL Server to Snowflake and Google Cloud BigQuery, cutting query costs by 35%.\n- Modeled dimensional schemas (Star, Snowflake) and implemented streaming ETL with Kafka.",
            "Data Pipeline Engineer | Insight Dynamics (2020 - 2022)\n- Wrote optimized complex SQL transformations and stored procedures in PostgreSQL and BigQuery.\n- Built automated data validation and ingestion frameworks using Python, Pandas, and Docker."
        ],
        "education": "Master of Science in Data Science, San Jose State University (2018 - 2020)\nBachelor of Technology in Computer Engineering (2014 - 2018)",
        "skills": "Python, SQL, Apache Spark, PySpark, Hadoop, Kafka, Airflow, Snowflake, BigQuery, Databricks, PostgreSQL, ETL, Docker, GCP, Git, Data Modeling",
        "formats": ["txt", "pdf"]
    },
    {
        "filename_base": "michael_brown_cybersecurity",
        "name": "Michael Brown",
        "title": "Cybersecurity & Incident Response Analyst",
        "contact": "michael.brown@secops-pro.com | (555) 654-7890 | Arlington, VA",
        "links": "LinkedIn: linkedin.com/in/michaelbrown-sec | GitHub: github.com/michaelbrown-cyber",
        "summary": "Dedicated Cybersecurity Specialist with 4+ years in threat hunting, vulnerability management, SIEM log monitoring, and network intrusion prevention. Certified CISSP and CEH.",
        "experience": [
            "Cybersecurity Operations Analyst | DefenseTech Shield (2022 - Present)\n- Monitored security incident and event management (SIEM) telemetry using Splunk and Microsoft Sentinel across 10,000 endpoints.\n- Performed vulnerability scans, threat mitigation, and penetration tests using Kali Linux, Wireshark, and Metasploit.\n- Automated threat intelligence feeds and triage scripts using Python and Bash.",
            "Security Analyst | Horizon InfoSec (2020 - 2022)\n- Analyzed network firewall logs, IDS/IPS rules, and investigated phishing incidents.\n- Implemented multi-factor authentication (MFA) and zero-trust identity policies in Azure Active Directory."
        ],
        "education": "Bachelor of Science in Cybersecurity & Information Assurance, George Mason University (2016 - 2020)",
        "skills": "Python, Bash, Linux, Network Security, Splunk, SIEM, Wireshark, Incident Response, Firewalls, Vulnerability Management, Azure, Critical Thinking",
        "formats": ["txt"]
    },
    {
        "filename_base": "aisha_patel_frontend_react",
        "name": "Aisha Patel",
        "title": "Senior Frontend & UI/UX Engineer",
        "contact": "aisha.patel@frontend-studio.com | (555) 876-5432 | Los Angeles, CA",
        "links": "LinkedIn: linkedin.com/in/aishapatel-ui | GitHub: github.com/aishapatel",
        "summary": "Passionate Frontend Web Engineer with 4+ years creating accessible, responsive, and blazing-fast web experiences. Master of React, Next.js, TypeScript, Tailwind CSS, Redux, and modern browser APIs.",
        "experience": [
            "Senior Frontend Developer | NextGen Web (2022 - Present)\n- Built modern server-side rendered web applications using React, Next.js, and TypeScript, improving SEO score to 98%.\n- Designed reusable component library utilizing Tailwind CSS, Storybook, and Figma tokens.\n- Integrated state management with Redux Toolkit and handled real-time communication via WebSockets.\n- Reduced client bundle size by 45% through dynamic imports, code splitting, and asset optimization.",
            "Frontend Developer | Creative Pixel (2020 - 2022)\n- Built interactive user interfaces with React, JavaScript, HTML5, and CSS3.\n- Wrote automated unit tests with Jest and React Testing Library in an Agile Scrum sprint setup."
        ],
        "education": "Bachelor of Science in Computer Science, University of Southern California (2016 - 2020)",
        "skills": "React, Next.js, TypeScript, JavaScript, HTML5, CSS3, Tailwind CSS, Redux, GraphQL, REST API, Jest, WebSockets, Git, UI/UX, Agile",
        "formats": ["txt", "docx"]
    },
    {
        "filename_base": "carlos_rodriguez_cloud_architect",
        "name": "Carlos Rodriguez",
        "title": "Principal Cloud Infrastructure Architect",
        "contact": "carlos.rodriguez@cloudarch.io | (555) 432-1098 | Denver, CO",
        "links": "LinkedIn: linkedin.com/in/carlosrodriguez-cloud | GitHub: github.com/carloscloud",
        "summary": "AWS Certified Solutions Architect with 7+ years of experience designing multi-region cloud infrastructures, Kubernetes container ecosystems, and automated disaster recovery systems.",
        "experience": [
            "Principal Cloud Architect | Apex Cloud Solutions (2021 - Present)\n- Architected scalable multi-account AWS and Azure enterprise cloud landing zones serving 20M users.\n- Standardized Infrastructure as Code (IaC) across 40 teams using Terraform, Terragrunt, and Ansible.\n- Implemented enterprise Kubernetes (EKS) clusters with service mesh (Istio), Prometheus, and Grafana.\n- Spearheaded cloud cost optimization program, delivering \$1.8M in annual cloud infrastructure savings.",
            "Senior DevOps Engineer | CloudScale Networks (2017 - 2021)\n- Built automated CI/CD deployment pipelines using GitHub Actions, Jenkins, and Helm.\n- Hardened Linux OS environments and managed automated security vulnerability patching."
        ],
        "education": "Master of Science in Information Systems, University of Colorado Boulder (2015 - 2017)\nBachelor of Science in Computer Science (2011 - 2015)",
        "skills": "AWS, Azure, GCP, Kubernetes, Docker, Terraform, Ansible, CI/CD, Prometheus, Grafana, Linux, Python, Bash, System Design, Leadership",
        "formats": ["txt", "pdf"]
    },
    {
        "filename_base": "rachel_green_mobile_dev",
        "name": "Rachel Green",
        "title": "Senior Cross-Platform Mobile Engineer",
        "contact": "rachel.green@appcraft.dev | (555) 321-6547 | Atlanta, GA",
        "links": "LinkedIn: linkedin.com/in/rachelgreen-mobile | GitHub: github.com/rachelgreen",
        "summary": "Mobile App Developer with 4+ years of experience publishing top-rated iOS and Android apps on Google Play and Apple App Store. Expert in Flutter, React Native, Dart, Swift, Kotlin, and Firebase.",
        "experience": [
            "Senior Mobile Developer | AppVenture Studios (2022 - Present)\n- Developed and published 3 cross-platform mobile apps using Flutter and Dart, garnering 500,000+ active downloads.\n- Integrated native plugins in Swift (iOS) and Kotlin (Android) for camera hardware and push notifications.\n- Built offline-first data caching with SQLite and synchronized backend data through GraphQL and REST APIs.\n- Configured automated mobile CI/CD pipelines using Fastlane and GitHub Actions.",
            "Mobile App Engineer | MobilityTech (2020 - 2022)\n- Engineered responsive client interfaces using React Native, TypeScript, and Redux.\n- Integrated Stripe in-app payments, Firebase analytics, and crash reporting."
        ],
        "education": "Bachelor of Science in Computer Science, Georgia Institute of Technology (2016 - 2020)",
        "skills": "Flutter, React Native, Dart, Swift, Kotlin, iOS, Android, REST API, GraphQL, Firebase, Git, Unit Testing, Mobile Architecture",
        "formats": ["txt"]
    },
    {
        "filename_base": "thomas_clark_bi_analyst",
        "name": "Thomas Clark",
        "title": "Senior Business Intelligence & Data Analyst",
        "contact": "thomas.clark@analytics-hub.com | (555) 210-9876 | Chicago, IL",
        "links": "LinkedIn: linkedin.com/in/thomasclark-bi | GitHub: github.com/thomasclark",
        "summary": "Data Analyst with 4+ years translating raw business data into actionable executive dashboards, revenue forecasting models, and automated KPI reporting. Proficient in SQL, Power BI, Tableau, Excel, and Python.",
        "experience": [
            "Lead BI Analyst | Strategic Growth Partners (2022 - Present)\n- Developed 15+ automated executive dashboards in Power BI and Tableau connected to PostgreSQL and Snowflake.\n- Authored advanced DAX measures and optimized complex SQL analytical queries, decreasing dashboard refresh time by 70%.\n- Built statistical regression and forecasting models in Python (Pandas, SciPy, Matplotlib) to predict customer retention.",
            "Data Analyst | Global Metrics (2020 - 2022)\n- Conducted A/B testing and marketing attribution analysis for e-commerce campaigns.\n- Automated weekly financial reporting workflows using Excel (VBA, Power Query) and Python scripts."
        ],
        "education": "Bachelor of Science in Statistics & Economics, University of Illinois (2016 - 2020)",
        "skills": "SQL, Power BI, Tableau, Excel, Python, Pandas, Matplotlib, Statistical Modeling, PostgreSQL, Snowflake, Communication, Problem Solving",
        "formats": ["txt"]
    }
]


def generate_txt(res: dict):
    filepath = os.path.join(OUTPUT_DIR, f"{res['filename_base']}.txt")
    lines = [
        res["name"],
        f"Title: {res['title']}",
        f"Contact: {res['contact']}",
        res["links"],
        "\nPROFESSIONAL SUMMARY",
        res["summary"],
        "\nWORK EXPERIENCE",
        "\n\n".join(res["experience"]),
        "\nEDUCATION",
        res["education"],
        "\nTECHNICAL SKILLS",
        res["skills"]
    ]
    with open(filepath, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Generated TXT: {filepath}")


def generate_docx(res: dict):
    filepath = os.path.join(OUTPUT_DIR, f"{res['filename_base']}.docx")
    doc = docx.Document()
    doc.add_heading(res["name"], 0)
    doc.add_paragraph(f"{res['title']} | {res['contact']}")
    doc.add_paragraph(res["links"])
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(res["summary"])
    
    doc.add_heading("Work Experience", level=1)
    for exp in res["experience"]:
        doc.add_paragraph(exp)
        
    doc.add_heading("Education", level=1)
    doc.add_paragraph(res["education"])
    
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(res["skills"])
    
    doc.save(filepath)
    print(f"Generated DOCX: {filepath}")


def generate_pdf(res: dict):
    filepath = os.path.join(OUTPUT_DIR, f"{res['filename_base']}.pdf")
    doc = SimpleDocTemplate(filepath, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('NameTitle', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor('#1E3A8A'))
    subtitle_style = ParagraphStyle('ContactSub', parent=styles['Normal'], fontSize=9, leading=12, textColor=colors.HexColor('#475569'))
    heading_style = ParagraphStyle('SectionHeading', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor('#2563EB'), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=9, leading=13, textColor=colors.HexColor('#1E293B'))
    
    story = [
        Paragraph(res["name"], title_style),
        Paragraph(f"<b>{res['title']}</b> | {res['contact']}", subtitle_style),
        Paragraph(res["links"], subtitle_style),
        Spacer(1, 8),
        Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style),
        Paragraph(res["summary"], body_style),
        Spacer(1, 6),
        Paragraph("<b>WORK EXPERIENCE</b>", heading_style)
    ]
    
    for exp in res["experience"]:
        for line in exp.split("\n"):
            story.append(Paragraph(line, body_style))
        story.append(Spacer(1, 4))
        
    story.append(Paragraph("<b>EDUCATION</b>", heading_style))
    for line in res["education"].split("\n"):
        story.append(Paragraph(line, body_style))
        
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>TECHNICAL SKILLS</b>", heading_style))
    story.append(Paragraph(res["skills"], body_style))
    
    doc.build(story)
    print(f"Generated PDF: {filepath}")


def main():
    for res in RESUMES_DATA:
        generate_txt(res)
        if "docx" in res["formats"]:
            generate_docx(res)
        if "pdf" in res["formats"]:
            generate_pdf(res)
    print(f"\nAll sample resumes created in {OUTPUT_DIR}!")

if __name__ == "__main__":
    main()
