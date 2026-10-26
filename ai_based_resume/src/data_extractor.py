"""
Data Extractor Module
Handles extraction of raw and normalized text from PDF, DOCX, and TXT files.
Supports file paths as well as in-memory file-like objects (e.g. Streamlit UploadedFile).
"""

import io
import re
from typing import Union, BinaryIO
import pypdf
import docx


def clean_extracted_text(text: str) -> str:
    """
    Cleans extracted raw text from PDF/DOCX to remove extraction artifacts.
    """
    if not text:
        return ""
    
    # Replace non-breaking spaces and carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("\u200b", "")
    
    # Merge broken hyphenated words at line breaks (e.g., "engi-\nneer" -> "engineer")
    text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)
    
    # Replace multiple consecutive spaces/tabs with a single space
    text = re.sub(r'[ \t]+', ' ', text)
    
    # Collapse 3 or more newlines into double newlines
    text = re.sub(r'\n{3,}', '\n\n', text)
    
    return text.strip()


def extract_text_from_pdf(file_source: Union[str, BinaryIO, io.BytesIO]) -> str:
    """
    Extracts text from a PDF file using pypdf.
    """
    try:
        reader = pypdf.PdfReader(file_source)
        text_parts = []
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return clean_extracted_text("\n".join(text_parts))
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_source: Union[str, BinaryIO, io.BytesIO]) -> str:
    """
    Extracts text from a DOCX file using python-docx.
    """
    try:
        doc = docx.Document(file_source)
        full_text = []
        # Extract paragraph texts
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        # Extract table texts
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return clean_extracted_text("\n".join(full_text))
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_txt(file_source: Union[str, BinaryIO, io.BytesIO]) -> str:
    """
    Extracts text from a TXT file with fallback encodings.
    """
    try:
        if isinstance(file_source, str):
            with open(file_source, "r", encoding="utf-8", errors="ignore") as f:
                return clean_extracted_text(f.read())
        else:
            if hasattr(file_source, "seek"):
                file_source.seek(0)
            content = file_source.read()
            if isinstance(content, bytes):
                try:
                    text = content.decode("utf-8")
                except UnicodeDecodeError:
                    text = content.decode("latin-1", errors="ignore")
            else:
                text = str(content)
            return clean_extracted_text(text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from TXT: {str(e)}")


def extract_text(file_source: Union[str, BinaryIO, io.BytesIO], filename: str = "") -> str:
    """
    Universal text extractor detecting file format by filename or path.
    """
    target_name = ""
    if isinstance(file_source, str):
        target_name = file_source.lower()
    elif filename:
        target_name = filename.lower()
    elif hasattr(file_source, "name"):
        target_name = getattr(file_source, "name", "").lower()
    
    if target_name.endswith(".pdf"):
        return extract_text_from_pdf(file_source)
    elif target_name.endswith(".docx"):
        return extract_text_from_docx(file_source)
    elif target_name.endswith(".txt"):
        return extract_text_from_txt(file_source)
    else:
        # Default try PDF first, then fallback to text
        try:
            return extract_text_from_pdf(file_source)
        except Exception:
            return extract_text_from_txt(file_source)
